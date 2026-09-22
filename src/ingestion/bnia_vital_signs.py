from openpyxl import load_workbook

from src.database.load_bnia import load_bnia_data


BNIA_FILEPATH = "data/raw/bnia_vital_signs/Vital Signs 2010-2023.xlsx"


def load_bnia_workbook():
    return load_workbook(
        BNIA_FILEPATH,
        read_only=True,
        data_only=True,
    )


def read_bnia_year(workbook, year):
    worksheet = workbook[str(year)]

    rows = list(worksheet.iter_rows(values_only=True))

    indicator_numbers = rows[0]
    headers = rows[1]
    data_rows = rows[2:]

    return indicator_numbers, headers, data_rows


def validate_bnia_year(indicator_numbers, headers, data_rows):
    if len(indicator_numbers) != len(headers):
        raise ValueError("BNIA indicator numbers and headers have different lengths")

    if len(data_rows) != 56:
        raise ValueError(f"Unexpected BNIA data row count: {len(data_rows)}")

    if headers[1] != "CSA2010":
        raise ValueError(f"Unexpected BNIA geography header: {headers[1]}")

    for row in data_rows:
        if len(row) != len(headers):
            raise ValueError(f"Unexpected BNIA row length: {len(row)}")


def clean_bnia_value(value):
    if isinstance(value, (int, float)):
        return value

    return None


def transform_bnia_year(year, indicator_numbers, headers, data_rows):
    records = []

    for row in data_rows:
        csa = row[1]

        for column_index in range(2, len(headers)):
            column_values = [row[column_index] for row in data_rows]

            if all(value is None or value == "--" for value in column_values):
                continue

            indicator = clean_bnia_header(headers[column_index])
            indicator_number = indicator_numbers[column_index]
            raw_value = row[column_index]
            value = clean_bnia_value(raw_value)

            if indicator is None or indicator == "CSA2020":
                continue

            record = {
                "source": "BNIA Vital Signs",
                "year": int(year),
                "csa": csa,
                "indicator_number": indicator_number,
                "indicator": indicator,
                "raw_value": raw_value,
                "value": value
            }

            records.append(record)

    return records


def clean_bnia_header(header):
    if isinstance(header, str):
        return header.strip()

    return header


def transform_bnia_workbook(workbook):
    records = []

    for year in workbook.sheetnames:
        indicator_numbers, headers, data_rows = read_bnia_year(workbook, year)

        validate_bnia_year(
            indicator_numbers,
            headers,
            data_rows,
        )

        year_records = transform_bnia_year(
            year,
            indicator_numbers,
            headers,
            data_rows,
        )

        records.extend(year_records)

    validate_unique_bnia_records(records)

    return records


def validate_unique_bnia_records(records):
    keys = [
        (
            record["year"],
            record["csa"],
            record["indicator_number"],
            record["indicator"],
        )
        for record in records
    ]

    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate BNIA records found")


def transform_bnia_geography(year, headers, data_rows):
    geography_records = []

    csa2010_index = headers.index("CSA2010")

    if "CSA2020" in headers:
        csa2020_index = headers.index("CSA2020")
    else:
        csa2020_index = None

    for row in data_rows:
        csa2010 = row[csa2010_index]

        if csa2020_index is not None:
            csa2020 = row[csa2020_index]
        else:
            csa2020 = None

        geography_record = {
            "source": "BNIA Vital Signs",
            "year": int(year),
            "csa2010": csa2010,
            "csa2020": csa2020,
            "is_citywide": csa2010 == "Baltimore City",
        }

        geography_records.append(geography_record)

    return geography_records


def transform_bnia_geographies(workbook):
    geography_records = []

    for year in workbook.sheetnames:
        indicator_numbers, headers, data_rows = read_bnia_year(workbook, year)

        year_geographies = transform_bnia_geography(
            year,
            headers,
            data_rows,
        )

        geography_records.extend(year_geographies)

    return geography_records


def validate_unique_bnia_geographies(records):
    keys = [
        (
            record["year"],
            record["csa2010"],
        )
        for record in records
    ]

    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate BNIA geography records found")


def main():
    workbook = load_bnia_workbook()

    try:
        observations = transform_bnia_workbook(workbook)
        geographies = transform_bnia_geographies(workbook)

        print("BNIA geographies:", len(geographies))
        print("BNIA observations:", len(observations))

        load_bnia_data(geographies, observations)

        print("BNIA load complete")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()